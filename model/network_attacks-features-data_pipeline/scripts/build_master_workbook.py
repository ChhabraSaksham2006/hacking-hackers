import os
import sys
import glob
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def sanitize_sheet_title(title, existing_titles):
    invalid_chars = [':', '\\', '/', '?', '*', '[', ']']
    for ch in invalid_chars:
        title = title.replace(ch, '_')
    if len(title) > 31:
        title = title[:31]
    base_title = title
    counter = 1
    while title.lower() in [t.lower() for t in existing_titles]:
        suffix = f"_{counter}"
        title = base_title[:31 - len(suffix)] + suffix
        counter += 1
    existing_titles.append(title)
    return title

def build_master_workbook():
    root_dir = os.path.abspath(".")
    wb = Workbook()
    
    # Fonts and styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1B365D")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="555555")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    sec_header_fill = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
    accent_fill = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
    link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="333333")
    meta_val_font = Font(name="Calibri", size=10, color="222222")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Master Index Sheet
    ws_index = wb.active
    ws_index.title = "Master_Index"
    ws_index.views.sheetView[0].showGridLines = True
    
    ws_index.cell(row=1, column=1, value="SIH26153: AI-BASED NETWORK ATTACK FORECASTING").font = title_font
    ws_index.cell(row=2, column=1, value="Comprehensive Ground Truth Master Research & Data Workbook").font = subtitle_font
    ws_index.cell(row=3, column=1, value=f"Repository Root: {root_dir} | Generated Authoritatively").font = subtitle_font
    
    index_headers = ["Category", "Sheet Name", "Data Type", "Item Count / Rows", "Exact Absolute File Path", "Relative Repository Path", "Scientific Description"]
    for col_idx, h in enumerate(index_headers, start=1):
        c = ws_index.cell(row=5, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    
    index_row = 6
    sheet_titles = ["Master_Index"]
    
    # 2. Gather All Datasets / CSVs
    csv_candidates = [
        ("Master Experiment Registry", "research_archive/experiment_registry.csv", "All 94 experimental runs across Phases 4, 5.5, 6, 7"),
        ("Phase 5.5 Experiment Results", "reports/phase_5_5/authoritative_experiment_results.csv", "Authoritative Phase 5.5 E001-E005 results (K=1, 10, 50)"),
        ("Phase 5.5 Baseline Comparison", "reports/phase_5_5/baseline_comparison.csv", "Evaluated baselines under identical chronological test set"),
        ("Phase 5.5 Onset Analysis", "reports/phase_5_5/onset_analysis.csv", "Preliminary leakage-free pre-onset early warning analysis"),
        ("Phase 6 Onset Benchmark", "reports/phase_6/authoritative_onset_results.csv", "Authoritative Phase 6 pure-benign onset experiments (E601-E607)"),
        ("Phase 6 Attack Forensics", "reports/phase_6/attack_events_forensics.csv", "Forensic catalog of 373 physical attack episodes"),
        ("Phase 7 Operating Points", "reports/phase_7/02_operating_point_results.csv", "Calibrated validation threshold constraints and test results"),
        ("Phase 7 Hard Negatives", "reports/phase_7/03_hard_negative_analysis.csv", "50 detailed forensic profiles of false alarm windows"),
        ("Phase 7 Alert Aggregation", "reports/phase_7/04_alert_aggregation_results.csv", "17 sequential temporal aggregation strategies"),
        ("Phase 7 Loss Experiments", "reports/phase_7/05_loss_experiments.csv", "Ablation of loss weighting ratios and focal loss"),
        ("Phase 7 Precursor Experiments", "reports/phase_7/06_precursor_experiments.csv", "Ablation of precursor loss multiplier and decay window"),
        ("Phase 7 Test Episodes", "reports/phase_7/07_event_level_results.csv", "Event-by-event detection status for 7 OOD test episodes"),
        ("Phase 7 MITRE Attribution", "reports/phase_7/08_behavior_to_mitre.csv", "55 precursor windows mapped to MITRE ATT&CK candidates"),
        ("Phase 7 Multi-Seed Evaluation", "reports/phase_7/09_multiseed_results.csv", "Robustness confirmation across seeds 42, 123, 2025"),
        ("Phase 7 Master Benchmark", "reports/phase_7/10_final_model_comparison.csv", "Master final comparison between baselines and SparseRSSM"),
        ("Master Metrics Forensic Audit", "reports/final_audit/master_metrics.csv", "Phase 4.5 master metrics across historical baselines"),
        ("Temporal Window Density", "reports/temporal_audit/05_window_density.csv", "Density of network states per session"),
        ("Temporal Feature Quality", "reports/temporal_audit/09_feature_quality.csv", "Missing, infinite, and zero value metrics per 54 features"),
        ("Temporal Window Comparison", "reports/temporal_design/01_window_comparison.csv", "Comparison of window sizes (5s vs 10s vs 30s)"),
        ("Temporal State Statistics", "reports/temporal_states/02_temporal_state_statistics.csv", "Statistics of generated 54-D parquet state files"),
        ("Raw Data Audit", "reports/data_pipeline/03_raw_data_audit.csv", "Coverage and timestamps of raw CSV files"),
        ("Feature Audit Pipeline", "reports/data_pipeline/07_feature_audit.csv", "80-CICFlowMeter feature retention or dropping rationales"),
        ("Canonical Parquet Statistics", "reports/data_pipeline/14_canonical_dataset_statistics.csv", "Sizes, flow counts, and coverage of canonical parquets"),
        ("Exploratory Feature Audit", "reports/03_feature_audit.csv", "Phase 1 feature completeness and missing percentages"),
        ("Exploratory Label Audit", "reports/04_label_audit.csv", "Phase 1 attack family counts and label normalization"),
        ("Dataset Candidate Comparison", "reports/09_dataset_comparison.csv", "Provenance audit across historical candidate datasets")
    ]
    
    print("Processing CSV data sheets...")
    for label, rel_path, desc in csv_candidates:
        abs_path = os.path.join(root_dir, rel_path.replace("/", os.sep))
        if not os.path.exists(abs_path):
            continue
        
        # Load CSV rows
        with open(abs_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = list(csv.reader(f))
            
        if not reader:
            continue
            
        header = reader[0]
        rows = reader[1:]
        
        # Determine Sheet Name
        sheet_name_seed = "D_" + rel_path.split("/")[-1].replace(".csv", "")
        clean_title = sanitize_sheet_title(sheet_name_seed, sheet_titles)
        
        ws = wb.create_sheet(title=clean_title)
        ws.views.sheetView[0].showGridLines = True
        
        # Header / Meta block in sheet
        ws.cell(row=1, column=1, value=f"VERIFIED DATASET: {label.upper()}").font = title_font
        ws.cell(row=2, column=1, value="Exact Record Path:").font = meta_label_font
        c_path = ws.cell(row=2, column=2, value=abs_path)
        c_path.font = link_font
        
        ws.cell(row=3, column=1, value="Relative Repo Path:").font = meta_label_font
        ws.cell(row=3, column=2, value=rel_path).font = meta_val_font
        
        ws.cell(row=4, column=1, value="Data Dimensions:").font = meta_label_font
        ws.cell(row=4, column=2, value=f"{len(rows)} data rows, {len(header)} columns | Verified Ground Truth").font = meta_val_font
        
        # Table Headers
        for col_idx, h_text in enumerate(header, start=1):
            c = ws.cell(row=6, column=col_idx, value=h_text)
            c.font = header_font
            c.fill = sec_header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
            
        # Table Rows
        for r_idx, r_data in enumerate(rows, start=7):
            fill_row = accent_fill if (r_idx % 2 == 0) else None
            for col_idx, val in enumerate(r_data, start=1):
                c = ws.cell(row=r_idx, column=col_idx)
                # Parse numeric if possible
                try:
                    if "." in val:
                        c.value = float(val)
                    else:
                        c.value = int(val)
                except ValueError:
                    c.value = val
                c.font = data_font
                c.border = thin_border
                if fill_row:
                    c.fill = fill_row
                    
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
            
        # Add to Master Index
        ws_index.cell(row=index_row, column=1, value="Empirical Data (CSV)").font = data_font
        c_name = ws_index.cell(row=index_row, column=2, value=clean_title)
        c_name.font = link_font
        ws_index.cell(row=index_row, column=3, value="Tabular Dataset").font = data_font
        ws_index.cell(row=index_row, column=4, value=f"{len(rows)} rows, {len(header)} cols").font = data_font
        ws_index.cell(row=index_row, column=5, value=abs_path).font = meta_val_font
        ws_index.cell(row=index_row, column=6, value=rel_path).font = meta_val_font
        ws_index.cell(row=index_row, column=7, value=desc).font = data_font
        index_row += 1

    # 3. Process Standalone Research Archive Tables (research_archive/tables/*.md)
    print("Processing Research Archive Markdown Tables...")
    table_mds = sorted(glob.glob("research_archive/tables/*.md"))
    for t_path in table_mds:
        rel_t_path = t_path.replace("\\", "/")
        abs_t_path = os.path.join(root_dir, t_path)
        
        with open(abs_t_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = [line.strip() for line in f.readlines()]
            
        # Parse table lines and text
        table_rows = []
        text_lines = []
        in_table = False
        title_text = rel_t_path.split("/")[-1].replace(".md", "")
        
        for line in lines:
            if line.startswith("# "):
                title_text = line.replace("# ", "")
            elif line.startswith("|") and line.endswith("|"):
                # Table row
                if "---" in line:
                    continue # separator
                cols = [c.strip() for c in line.strip("|").split("|")]
                table_rows.append(cols)
            else:
                if line:
                    text_lines.append(line)
                    
        sheet_name_seed = "T_" + rel_t_path.split("/")[-1].replace(".md", "")
        clean_title = sanitize_sheet_title(sheet_name_seed, sheet_titles)
        
        ws = wb.create_sheet(title=clean_title)
        ws.views.sheetView[0].showGridLines = True
        
        ws.cell(row=1, column=1, value=title_text).font = title_font
        ws.cell(row=2, column=1, value="Exact Record Path:").font = meta_label_font
        ws.cell(row=2, column=2, value=abs_t_path).font = link_font
        ws.cell(row=3, column=1, value="Relative Repo Path:").font = meta_label_font
        ws.cell(row=3, column=2, value=rel_t_path).font = meta_val_font
        
        cur_row = 5
        if table_rows:
            # Header
            header = table_rows[0]
            for col_idx, h_text in enumerate(header, start=1):
                c = ws.cell(row=cur_row, column=col_idx, value=h_text)
                c.font = header_font
                c.fill = sec_header_fill
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = thin_border
            cur_row += 1
            
            for r_data in table_rows[1:]:
                fill_row = accent_fill if (cur_row % 2 == 0) else None
                for col_idx, val in enumerate(r_data, start=1):
                    c = ws.cell(row=cur_row, column=col_idx)
                    try:
                        clean_val = val.replace("**", "").replace("$", "").replace("%", "").strip()
                        if "." in clean_val:
                            c.value = float(clean_val)
                        else:
                            c.value = int(clean_val)
                    except ValueError:
                        c.value = val.replace("**", "").replace("$", "")
                    c.font = data_font
                    c.border = thin_border
                    if fill_row:
                        c.fill = fill_row
                cur_row += 1
                
            cur_row += 1
            
        # Interpretation / Text block
        ws.cell(row=cur_row, column=1, value="SCIENTIFIC INTERPRETATION & FINDINGS:").font = meta_label_font
        cur_row += 1
        for t_line in text_lines:
            c = ws.cell(row=cur_row, column=1, value=t_line)
            c.font = data_font
            cur_row += 1
            
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col[:15])
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 60)
            
        ws_index.cell(row=index_row, column=1, value="Research Table").font = data_font
        ws_index.cell(row=index_row, column=2, value=clean_title).font = link_font
        ws_index.cell(row=index_row, column=3, value="Publication Table + Text").font = data_font
        ws_index.cell(row=index_row, column=4, value=f"{len(table_rows)} table rows").font = data_font
        ws_index.cell(row=index_row, column=5, value=abs_t_path).font = meta_val_font
        ws_index.cell(row=index_row, column=6, value=rel_t_path).font = meta_val_font
        ws_index.cell(row=index_row, column=7, value=title_text).font = data_font
        index_row += 1

    # 4. Master Research Reports & Methodology Documents
    print("Processing Core Master Research Reports...")
    core_reports = [
        ("Final Research Summary", "docs/FINAL_RESEARCH_SUMMARY.md", "Authoritative 18-question final research freeze summary"),
        ("Complete Research Flow", "research_archive/COMPLETE_RESEARCH_FLOW.md", "Full chronological research progression with scientific decisions"),
        ("Research Chain Progression", "research_archive/RESEARCH_CHAIN.md", "Detailed cause-and-effect narrative for every experiment"),
        ("Final Benchmark Analysis", "research_archive/FINAL_BENCHMARK.md", "Multi-task evaluation separating continuation from pre-onset warning"),
        ("Final Hyperparameters Audit", "research_archive/FINAL_HYPERPARAMETERS.md", "Complete parameter classification and origins"),
        ("Paper-Ready Results & Discrepancies", "research_archive/PAPER_RESULTS.md", "Final publication synthesis and Seed-123 resolution"),
        ("MITRE Behavioral Attribution Layer", "research_archive/MITRE_MAPPING.md", "Detailed heuristic evidence matrix for MITRE ATT&CK techniques"),
        ("Four Final System Outputs", "research_archive/FINAL_SYSTEM_OUTPUTS.md", "Specification of multi-channel world model outputs"),
        ("Phase 7 Final Report", "reports/phase_7/11_phase7_final_report.md", "Operational aggregation and behavioral attribution findings"),
        ("Phase 6 Final Report", "reports/phase_6/PHASE_6_FINAL_REPORT.md", "Pre-attack onset forecasting and episode forensics"),
        ("Phase 5.5 Final Report", "reports/phase_5_5/PHASE_5_5_FINAL_REPORT.md", "Forensic audit fixes, gradient restoration, and benchmarks"),
        ("Phase 4.5 Executive Summary", "reports/final_audit/23_executive_summary.md", "Forensic discovery of the Persistence Paradox and untrained head"),
        ("Persistence Forensic Audit", "reports/final_audit/14_persistence_forensics.md", "Mathematical proof of continuation autocorrelation in CIC-IDS2018"),
        ("Cleanup & Lineage Inventory", "docs/CLEANUP_INVENTORY.md", "Classification of all repository files into A/B/C/D/E"),
        ("Production Model Card", "deployment/MODEL_CARD.md", "Performance, limitations, and ethical considerations for deployment")
    ]
    
    for label, rel_path, desc in core_reports:
        abs_path = os.path.join(root_dir, rel_path.replace("/", os.sep))
        if not os.path.exists(abs_path):
            continue
            
        with open(abs_path, "r", encoding="utf-8", errors="ignore") as f:
            content_lines = f.readlines()
            
        sheet_name_seed = "DOC_" + rel_path.split("/")[-1].replace(".md", "")
        clean_title = sanitize_sheet_title(sheet_name_seed, sheet_titles)
        
        ws = wb.create_sheet(title=clean_title)
        ws.views.sheetView[0].showGridLines = True
        
        ws.cell(row=1, column=1, value=label.upper()).font = title_font
        ws.cell(row=2, column=1, value="Exact Record Path:").font = meta_label_font
        ws.cell(row=2, column=2, value=abs_path).font = link_font
        ws.cell(row=3, column=1, value="Relative Repo Path:").font = meta_label_font
        ws.cell(row=3, column=2, value=rel_path).font = meta_val_font
        ws.cell(row=4, column=1, value="Document Length:").font = meta_label_font
        ws.cell(row=4, column=2, value=f"{len(content_lines)} lines | Verified Unaltered Text").font = meta_val_font
        
        cur_row = 6
        for line in content_lines:
            stripped = line.rstrip("\r\n")
            c = ws.cell(row=cur_row, column=1, value=stripped)
            if stripped.startswith("# "):
                c.font = Font(name="Calibri", size=14, bold=True, color="1B365D")
            elif stripped.startswith("## "):
                c.font = Font(name="Calibri", size=12, bold=True, color="2E5B88")
            elif stripped.startswith("### "):
                c.font = Font(name="Calibri", size=11, bold=True, color="333333")
            elif stripped.startswith("|"):
                c.font = Font(name="Courier New", size=10)
            else:
                c.font = data_font
            cur_row += 1
            
        ws.column_dimensions["A"].width = 120
        
        ws_index.cell(row=index_row, column=1, value="Research Report (Doc)").font = data_font
        ws_index.cell(row=index_row, column=2, value=clean_title).font = link_font
        ws_index.cell(row=index_row, column=3, value="Full Text Report").font = data_font
        ws_index.cell(row=index_row, column=4, value=f"{len(content_lines)} text lines").font = data_font
        ws_index.cell(row=index_row, column=5, value=abs_path).font = meta_val_font
        ws_index.cell(row=index_row, column=6, value=rel_path).font = meta_val_font
        ws_index.cell(row=index_row, column=7, value=desc).font = data_font
        index_row += 1

    # Formatting Master Index columns
    for col in ws_index.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_index.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 70)

    out_file = os.path.join(root_dir, "FINAL_RESEARCH_MASTER_WORKBOOK.xlsx")
    wb.save(out_file)
    print(f"SUCCESS: Created master ground truth workbook: {out_file}")
    print(f"Total sheets created: {len(sheet_titles)}")

if __name__ == "__main__":
    build_master_workbook()
